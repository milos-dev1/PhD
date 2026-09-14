"""Word document utilities for B5 thesis (11pt) with hyperlinks, bookmarks, and page numbers."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from citations import inject_citations

# ISO B5
B5_WIDTH = Mm(176)
B5_HEIGHT = Mm(250)
BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(11)
HEADING_COLOR = RGBColor(0, 0, 0)
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)


def slugify(text: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "_", text.strip()).strip("_").lower()
    return s[:60] or "section"


def bookmark_name_for_heading(heading: str) -> str:
    return f"sec_{slugify(heading)}"


def bookmark_name_for_figure(caption: str) -> str | None:
    m = re.search(r"Figure\s+(\d+)\.(\d+)", caption, re.I)
    if m:
        return f"fig_{m.group(1)}_{m.group(2)}"
    return None


def bookmark_name_for_table(caption: str) -> str | None:
    m = re.search(r"Table\s+(\d+)\.(\d+)", caption, re.I)
    if m:
        return f"tbl_{m.group(1)}_{m.group(2)}"
    return None


_BOOKMARK_ID = 0


def _next_bookmark_id() -> str:
    global _BOOKMARK_ID
    _BOOKMARK_ID += 1
    return str(_BOOKMARK_ID)


def _add_bookmark(paragraph, name: str) -> None:
    """Add a Word bookmark covering the paragraph."""
    bid = _next_bookmark_id()
    tag = paragraph._p
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    tag.insert(0, start)
    tag.append(end)


def _add_hyperlink(paragraph, bookmark: str, text: str, *, bold: bool = False, size: Pt | None = None) -> None:
    """Add an internal hyperlink (bookmark jump) to a paragraph."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), bookmark)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), BODY_FONT)
    rFonts.set(qn("w:hAnsi"), BODY_FONT)
    rPr.append(rFonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int((size or BODY_SIZE).pt * 2)))
    rPr.append(sz)

    if bold:
        rPr.append(OxmlElement("w:b"))

    new_run.append(rPr)
    text_elem = OxmlElement("w:t")
    text_elem.set(qn("xml:space"), "preserve")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _add_run(paragraph, text: str, *, bold: bool = False, italic: bool = False, size: Pt | None = None) -> None:
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = size or BODY_SIZE
    run.bold = bold
    run.italic = italic


def add_paragraph_with_citations(
    doc: Document,
    text: str,
    *,
    bold: bool = False,
    indent: bool = False,
    auto_cite: bool = True,
) -> None:
    """Write a paragraph; convert [n] markers into hyperlinks to ref bookmarks."""
    if auto_cite:
        text = inject_citations(text)
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Mm(10)

    # Split on citation tokens like [1], [1,2], [21]
    parts = re.split(r"(\[\d+(?:\s*,\s*\d+)*\])", text)
    for part in parts:
        if not part:
            continue
        m = re.fullmatch(r"\[(\d+(?:\s*,\s*\d+)*)\]", part)
        if m:
            nums = [int(x.strip()) for x in m.group(1).split(",")]
            for i, n in enumerate(nums):
                if i > 0:
                    _add_run(p, ",", size=BODY_SIZE)
                _add_hyperlink(p, f"ref_{n}", f"[{n}]")
        else:
            _add_run(p, part, bold=bold)


def add_heading_bookmarked(doc: Document, text: str, level: int = 1) -> str:
    """Add a heading and bookmark it; return bookmark name."""
    heading = doc.add_heading(text, level=level)
    name = bookmark_name_for_heading(text)
    _add_bookmark(heading, name)
    return name


def add_page_numbers(doc: Document) -> None:
    """Add centered page numbers in the footer (PAGE field)."""
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Clear existing
    for r in list(p.runs):
        r.clear()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)
    for r in p.runs:
        r.font.name = BODY_FONT
        r.font.size = Pt(10)


def add_clickable_toc(doc: Document, entries: list[tuple[str, str | None, str]]) -> None:
    """
    Clickable table of contents.
    entries: list of (display_title, subtitle_or_None, bookmark_name)
    """
    h = doc.add_heading("Table of Contents", level=1)
    _add_bookmark(h, "sec_table_of_contents")
    for title, subtitle, bookmark in entries:
        p = doc.add_paragraph()
        indented = title.startswith("  ")
        display = title.strip()
        if indented:
            p.paragraph_format.left_indent = Mm(8)
        bold = display in (
            "Abstract",
            "Preface",
            "Conclusions",
            "Appendix",
            "References",
        ) or display.startswith("Chapter")
        _add_hyperlink(p, bookmark, display, bold=bold, size=Pt(10) if indented else BODY_SIZE)
        if subtitle:
            p2 = doc.add_paragraph()
            p2.paragraph_format.left_indent = Mm(8)
            _add_run(p2, subtitle, italic=True, size=Pt(10))
    doc.add_page_break()


def add_clickable_list(doc: Document, title: str, items: list[tuple[str, str]]) -> None:
    """Clickable list of figures/tables: (caption, bookmark)."""
    h = doc.add_heading(title, level=1)
    _add_bookmark(h, bookmark_name_for_heading(title))
    for caption, bookmark in items:
        p = doc.add_paragraph()
        _add_hyperlink(p, bookmark, caption, size=BODY_SIZE)
    doc.add_page_break()


def add_references(doc: Document, references: list[tuple]) -> None:
    """Numbered bibliography with bookmarks ref_1, ref_2, ... for citation links."""
    h = doc.add_heading("References", level=1)
    _add_bookmark(h, "sec_references")
    for i, ref in enumerate(references, 1):
        authors, title, venue, year = ref[:4]
        extra = ref[4] if len(ref) > 4 else ""
        text = f"[{i}] {authors}. \"{title}.\" {venue}, {year}."
        if extra:
            text += f" {extra}"
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Mm(10)
        p.paragraph_format.first_line_indent = Mm(-10)
        p.paragraph_format.space_after = Pt(4)
        _add_bookmark(p, f"ref_{i}")
        _add_run(p, text)


def create_thesis_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = B5_WIDTH
    section.page_height = B5_HEIGHT
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    for level, size in [(1, 16), (2, 14), (3, 12)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = HEADING_COLOR
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    add_page_numbers(doc)
    return doc


def add_title_page(doc: Document, title: str, subtitle: str, degree: str) -> None:
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = BODY_FONT

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(12)
    r2.font.name = BODY_FONT

    for _ in range(8):
        doc.add_paragraph()
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(degree)
    r3.font.size = Pt(12)
    r3.font.name = BODY_FONT
    doc.add_page_break()


def add_paragraph(doc: Document, text: str, *, bold: bool = False, indent: bool = False) -> None:
    add_paragraph_with_citations(doc, text, bold=bold, indent=indent, auto_cite=True)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    text = inject_citations(text)
    parts = re.split(r"(\[\d+(?:\s*,\s*\d+)*\])", text)
    for part in parts:
        if not part:
            continue
        m = re.fullmatch(r"\[(\d+(?:\s*,\s*\d+)*)\]", part)
        if m:
            nums = [int(x.strip()) for x in m.group(1).split(",")]
            for i, n in enumerate(nums):
                if i > 0:
                    _add_run(p, ",")
                _add_hyperlink(p, f"ref_{n}", f"[{n}]")
        else:
            _add_run(p, part)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Number")
    for run in p.runs:
        run.font.name = BODY_FONT
        run.font.size = BODY_SIZE


def add_figure(doc: Document, image_path: Path | None, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if image_path and image_path.exists():
        run = p.add_run()
        run.add_picture(str(image_path), width=Mm(140))
    else:
        run = p.add_run(f"[Figure: {caption}]")
        run.italic = True
        run.font.size = Pt(10)

    try:
        cap = doc.add_paragraph(caption, style="Caption")
    except KeyError:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption)
        cr.font.size = Pt(10)
        cr.font.name = BODY_FONT
        cr.italic = True
    else:
        for r in cap.runs:
            r.font.name = BODY_FONT
            r.font.size = Pt(10)

    bm = bookmark_name_for_figure(caption)
    if bm:
        _add_bookmark(cap, bm)
    doc.add_paragraph()


def add_figure_placeholder(doc: Document, caption: str) -> None:
    add_figure(doc, None, caption)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], caption: str) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
                r.font.name = BODY_FONT
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = val
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.name = BODY_FONT
    try:
        cap = doc.add_paragraph(caption, style="Caption")
    except KeyError:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption)
        cr.font.size = Pt(10)
        cr.font.name = BODY_FONT
        cr.italic = True
    else:
        for r in cap.runs:
            r.font.name = BODY_FONT
            r.font.size = Pt(10)

    bm = bookmark_name_for_table(caption)
    if bm:
        _add_bookmark(cap, bm)
    doc.add_paragraph()


def save_document(doc: Document, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def update_word_fields(docx_path: Path) -> bool:
    """Update PAGE fields and any TOC fields via Microsoft Word (Windows)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return False
    path = str(docx_path.resolve())
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(path, ReadOnly=False)
        doc.Fields.Update()
        for toc in doc.TablesOfContents:
            toc.Update()
        for tof in doc.TablesOfFigures:
            tof.Update()
        doc.Repaginate()
        doc.Save()
        doc.Close()
        return True
    except Exception:
        return False
    finally:
        try:
            word.Quit()
        except Exception:
            pass


# Backwards-compatible aliases used by older generators
def add_static_toc(doc: Document, chapters: list[tuple]) -> None:
    entries = []
    for entry in chapters:
        title = entry[0]
        sub = entry[1] if len(entry) > 1 else None
        bm = bookmark_name_for_heading(title.strip())
        entries.append((title, sub, bm))
    add_clickable_toc(doc, entries)


def add_static_list(doc: Document, title: str, items: list[str]) -> None:
    clickable: list[tuple[str, str]] = []
    for item in items:
        bm = bookmark_name_for_figure(item) or bookmark_name_for_table(item) or bookmark_name_for_heading(item)
        clickable.append((item, bm))
    add_clickable_list(doc, title, clickable)


def add_table_of_contents(doc: Document, title: str = "Table of Contents") -> None:
    add_heading_bookmarked(doc, title, 1)
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)
    doc.add_page_break()
